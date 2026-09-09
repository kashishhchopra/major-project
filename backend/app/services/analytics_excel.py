"""Render the analytics dashboard as a one-click Excel workbook.

The spreadsheet counterpart to services/analytics_pdf.py -- same aggregate
queries (app/api/analytics.py), one sheet per section, via openpyxl (already
a dependency for itinerary document parsing, so this adds nothing new).
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_table(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append(row)
    for col in range(1, len(headers) + 1):
        width = max([len(str(headers[col - 1]))] + [len(str(r[col - 1])) for r in rows] + [8])
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, 40)


def render_analytics_excel(
    summary: dict, over_time: list[dict], by_type: list[dict],
    zone_risk: list[dict], severity: list[dict],
    police_performance: dict, tourist_activity: dict, incident_types: dict,
) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    _write_table(ws, ["Metric", "Value"], [
        ["Total Tourists", summary["total_tourists"]],
        ["Active Tourists", summary["active_tourists"]],
        ["SOS Active", summary["sos_active"]],
        ["Missing", summary["missing"]],
        ["Total Incidents", summary["total_incidents"]],
        ["Open Incidents", summary["open_incidents"]],
        ["Active Alerts", summary["active_alerts"]],
        ["Avg Safety Score", summary["avg_safety_score"]],
        ["Avg Response Time (s)", summary["avg_response_time_seconds"]],
        ["Total Zones", summary["total_zones"]],
    ])

    ws = wb.create_sheet("Incidents Over Time")
    _write_table(ws, ["Date", "Count"], [[r["date"], r["count"]] for r in over_time])

    ws = wb.create_sheet("Incident Types")
    _write_table(ws, ["Type", "Count"], [[r["type"], r["count"]] for r in incident_types["by_type"]])
    sos = incident_types["sos"]
    ws.append([])
    ws.append(["SOS Total", sos["total"]])
    ws.append(["SOS Active", sos["active"]])
    ws.append(["SOS Resolved", sos["resolved"]])

    ws = wb.create_sheet("Alerts by Type")
    _write_table(ws, ["Type", "Count"], [[r["type"], r["count"]] for r in by_type])

    ws = wb.create_sheet("Zone Risk")
    _write_table(ws, ["Zone", "Risk Level", "Crime Index", "Alert Count"],
                [[r["zone"], r["risk_level"], r["crime_index"], r["alert_count"]] for r in zone_risk])

    ws = wb.create_sheet("Severity Breakdown")
    _write_table(ws, ["Severity", "Count"], [[r["severity"], r["count"]] for r in severity])

    ws = wb.create_sheet("Police Performance")
    _write_table(ws, ["Station", "Cases Handled", "Pending", "Resolved",
                      "Avg Response (s)", "Avg Resolution (s)", "Transfers Received"],
                [[r["station"], r["cases_handled"], r["pending"], r["resolved"],
                  r["avg_response_time_seconds"], r["avg_resolution_time_seconds"],
                  r["transfers_received"]] for r in police_performance["stations"]])

    ws = wb.create_sheet("Tourist Activity")
    _write_table(ws, ["Metric", "Value"], [
        ["Total Tourists", tourist_activity["total_tourists"]],
        ["Active Tourists", tourist_activity["active_tourists"]],
        ["Domestic Tourists", tourist_activity["domestic_tourists"]],
        ["International Tourists", tourist_activity["international_tourists"]],
    ])
    ws.append([])
    ws.append(["Popular Destination", "Visits"])
    for r in tourist_activity["popular_destinations"]:
        ws.append([r["destination"], r["count"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
